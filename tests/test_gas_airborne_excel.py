import os
import tempfile
import unittest

from openpyxl import load_workbook

from excel_generator import generate_gas_airborne_particle_excel


class GasAirborneParticleExcelTest(unittest.TestCase):
    def test_generates_separate_chart_sheets_per_grade(self):
        records = [
            {
                'no': '1',
                'management_number': 'CA-01',
                'location': '충전 3실 (2505)',
                'grade': 'A',
                'particle_05': 13,
                'particle_50': 0,
                'judgement': '적합',
                'criteria_text': '허용기준',
                'performed_date': '2025.02.28',
            },
            {
                'no': '2',
                'management_number': 'CA-02',
                'location': '충전 4실 (2506)',
                'grade': 'B',
                'particle_05': 628,
                'particle_50': 14,
                'judgement': '부적합',
                'criteria_text': '허용기준',
                'performed_date': '2025.08.28',
            },
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, 'gas-airborne.xlsx')
            generate_gas_airborne_particle_excel(records, path)
            workbook = load_workbook(path, data_only=False)

        self.assertEqual(
            workbook.sheetnames,
            [
                '데이터',
                'Pivot 0.5µm Grade A',
                'Pivot 0.5µm Grade B',
                'Pivot 5.0µm Grade A',
                'Pivot 5.0µm Grade B',
            ],
        )
        data_sheet = workbook['데이터']
        self.assertEqual(data_sheet['A1'].value, '부유입자 측정 일지')
        self.assertEqual(data_sheet['B6'].value, 'CA-02')
        self.assertEqual(data_sheet['E6'].value, 628)
        self.assertEqual(data_sheet['F6'].value, 14)
        self.assertEqual(data_sheet['H6'].value, '2025.08.28')
        self.assertEqual(data_sheet['E6'].fill.fill_type, 'solid')
        self.assertEqual(data_sheet['F6'].fill.fill_type, 'solid')
        self.assertEqual(data_sheet['G6'].fill.fill_type, 'solid')
        for particle_size in ('0.5', '5.0'):
            for grade in ('A', 'B'):
                chart_sheet = workbook[f'Pivot {particle_size}µm Grade {grade}']
                headers = [
                    chart_sheet.cell(row=1, column=column).value
                    for column in range(1, chart_sheet.max_column + 1)
                ]
                grade_limits = [
                    header for header in headers
                    if isinstance(header, str) and 'Grade' in header and '기준 =' in header
                ]
                expected_limit = {
                    ('0.5', 'A'): 23,
                    ('0.5', 'B'): 627,
                    ('5.0', 'A'): 4,
                    ('5.0', 'B'): 13,
                }[(particle_size, grade)]

                self.assertEqual(len(chart_sheet._charts), 1)
                self.assertFalse(chart_sheet._charts[0].visible_cells_only)
                self.assertEqual(
                    {
                        chart_sheet.cell(row=row, column=1).value
                        for row in range(2, chart_sheet.max_row + 1)
                        if chart_sheet.cell(row=row, column=1).value is not None
                    },
                    {grade},
                )
                self.assertEqual(
                    grade_limits,
                    [f'{grade} Grade 경고기준 = {expected_limit}'],
                )
                limit_chart = chart_sheet._charts[0]._charts[1]
                self.assertEqual(type(limit_chart).__name__, 'LineChart')
                self.assertIsNone(limit_chart.y_axis.majorGridlines)
                self.assertEqual(limit_chart.y_axis.crossBetween, 'midCat')
                self.assertEqual(len(limit_chart.ser), 1)
                self.assertRegex(
                    limit_chart.ser[0].val.numRef.f,
                    r'\$[A-Z]+\$2:\$[A-Z]+\$3$',
                )

    def test_keeps_all_available_grades_when_latest_date_is_first_half(self):
        records = [
            {
                'management_number': f'CA-0{index}',
                'location': f'Grade {grade} Room',
                'grade': grade,
                'particle_05': 1,
                'particle_50': 0,
                'judgement': '적합',
                'performed_date': '2026.02.01' if grade in ('A', 'B') else '2025.08.01',
            }
            for index, grade in enumerate('ABCD', start=1)
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, 'gas-airborne-all-grades.xlsx')
            generate_gas_airborne_particle_excel(records, path)
            workbook = load_workbook(path, data_only=False)

        for particle_size in ('0.5', '5.0'):
            for grade in 'ABCD':
                self.assertIn(
                    f'Pivot {particle_size}µm Grade {grade}',
                    workbook.sheetnames,
                )


if __name__ == '__main__':
    unittest.main()
